import os

import pypandoc # pypandoc will use the PANDOC_PATH set by the main app
import shutil
from docx import Document as DocxDocument
# from openpyxl import load_workbook # Add back if handling xlsx
# from pptx import Presentation     # Add back if handling pptx


TEMP_DIR_NAME = "redax_processing_temp_webview" # Unique temp dir

CREDENTIAL_KEYWORDS = [
    "password", "pwd", "secret", "username", "user name", "login", "user id",
    "credential", "credentials", "authorization", "bearer token", "api key",
    "client secret", "token", "auth key", "private key", "secret key", "access key"
]

def ensure_temp_dir():
    if not os.path.exists(TEMP_DIR_NAME):
        os.makedirs(TEMP_DIR_NAME)

def cleanup_temp_dir():
    if os.path.exists(TEMP_DIR_NAME):
        try:
            shutil.rmtree(TEMP_DIR_NAME)
            print(f"INFO: Cleaned up temp directory: {TEMP_DIR_NAME}")
        except Exception as e:
            print(f"ERROR: Could not remove temp directory {TEMP_DIR_NAME}: {e}")


def _redact_text_content_logic(text, active_patterns_compiled, custom_keywords_list, redact_credential_lines_enabled=False):
    processed_text_intermediate = str(text)

    if redact_credential_lines_enabled:
        lines = processed_text_intermediate.splitlines()
        new_lines = []
        for line in lines:
            if any(keyword.lower() in line.lower() for keyword in CREDENTIAL_KEYWORDS):
                new_lines.append("[REDACTED LINE]")
            else:
                new_lines.append(line)
        processed_text_intermediate = "\n".join(new_lines)

    redacted_text = processed_text_intermediate
    for pattern_regex in active_patterns_compiled:
        # Skip placeholder for credential line redaction if it was somehow compiled (it shouldn't be)
        if pattern_regex.pattern == "KEYWORD_LINE_REDACTION_PLACEHOLDER":
            continue
        redacted_text = pattern_regex.sub("[REDACTED]", redacted_text)
    for keyword in custom_keywords_list:
        if keyword:
            redacted_text = redacted_text.replace(keyword, "[REDACTED]")
    return redacted_text

def process_document_for_redaction(original_filepath, active_patterns_compiled, custom_keywords_list, output_format="md", redact_credential_lines_enabled=False):
    ensure_temp_dir()
    original_filename = os.path.basename(original_filepath)
    base_name, file_ext = os.path.splitext(original_filename)
    file_ext = file_ext.lower()

    output_filename_in_temp = f"{base_name}_redacted.{output_format}"
    final_output_path_in_temp = os.path.join(TEMP_DIR_NAME, output_filename_in_temp)

    temp_input_for_pandoc = None
    content_for_pandoc = None

    try:
        if file_ext == ".docx":
            doc = DocxDocument(original_filepath)
            for para in doc.paragraphs:
                para.text = _redact_text_content_logic(para.text, active_patterns_compiled, custom_keywords_list, redact_credential_lines_enabled)
            # Add table redaction etc. if needed
            # TODO: Consider tables, headers, footers for DOCX if keyword line redaction is enabled
            temp_input_for_pandoc = os.path.join(TEMP_DIR_NAME, f"temp_{original_filename}")
            doc.save(temp_input_for_pandoc)
        elif file_ext in [".md", ".txt"]:
            with open(original_filepath, "r", encoding="utf-8", errors='ignore') as f:
                content = f.read()
            content_for_pandoc = _redact_text_content_logic(content, active_patterns_compiled, custom_keywords_list, redact_credential_lines_enabled)
        elif file_ext == ".rtf":
            # For RTF, convert to MD first, then redact the MD content
            md_content = pypandoc.convert_file(original_filepath, 'markdown_strict', format='rtf', extra_args=['--wrap=none'])
            content_for_pandoc = _redact_text_content_logic(md_content, active_patterns_compiled, custom_keywords_list, redact_credential_lines_enabled)
        # Add other file types (xlsx, pptx) similarly, preparing either temp_input_for_pandoc or content_for_pandoc
        # For xlsx/pptx, you might copy to temp_input_for_pandoc and let Pandoc extract text if direct redaction is too complex.
        else:
            return {"error": f"Unsupported file type: {original_filename}"}

        pandoc_extra_args = ['--standalone']
        if output_format == 'pdf':
            pandoc_extra_args.append('--toc')

        if temp_input_for_pandoc:
            if output_format == "md":
                pypandoc.convert_file(temp_input_for_pandoc, 'markdown_strict', format=file_ext.lstrip('.'),
                                      outputfile=final_output_path_in_temp, extra_args=pandoc_extra_args + ['--wrap=none'])
            elif output_format == "pdf":
                pypandoc.convert_file(temp_input_for_pandoc, 'pdf', format=file_ext.lstrip('.'),
                                      outputfile=final_output_path_in_temp, extra_args=pandoc_extra_args)
            os.remove(temp_input_for_pandoc)
        elif content_for_pandoc:
            pypandoc.convert_text(content_for_pandoc, output_format, format='markdown',
                                  outputfile=final_output_path_in_temp, extra_args=pandoc_extra_args)
        else:
            return {"error": f"No content to process for {original_filename}"}

        return {"original_name": original_filename, "output_path": final_output_path_in_temp, "output_format": output_format}

    except OSError as e_pandoc_os_error: # MODIFIED to catch OSError for Pandoc issues
        # Check if the error message indicates Pandoc is missing or not executable
        error_str = str(e_pandoc_os_error).lower()
        if 'pandoc' in error_str and ('not found' in error_str or 'no such file' in error_str or 'permission denied' in error_str or 'isn\'t executable' in error_str):
            print(f"CRITICAL ERROR: Pandoc not found or not executable: {e_pandoc_os_error}")
            return {"error": f"Pandoc not found/executable. Please ensure Pandoc is correctly configured/bundled. Details: {e_pandoc_os_error}"}
        else:
            # Handle other OSErrors during Pandoc processing
            print(f"An OS Error occurred during Pandoc processing {original_filename}: {e_pandoc_os_error}")
            return {"error": f"OS error during processing {original_filename}: {str(e_pandoc_os_error)}"}
    except Exception as e: # General catch-all for other errors
        print(f"Error processing {original_filename}: {e}")
        return {"error": f"Error processing {original_filename}: {str(e)}"}
